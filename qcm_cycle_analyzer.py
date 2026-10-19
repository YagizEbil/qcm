#!/usr/bin/env python3

OUTPUT_FOLDER = None

FREQ_CHANNEL = "auto"

DELTA_F_SIGN = "f0_minus_f"

HAND_SELECTION = True
SNAP_TO_JUMP = True
SNAP_WINDOW_S = 4.0
REUSE_SAVED_SELECTION = True

EVENT_LABELING = "auto"

SENSITIVITY = 6.0

MIN_JUMP_HZ = 0.8

MERGE_WINDOW_S = 8.0

BASELINE_WINDOW_S = 5.0
END_WINDOW_S = 5.0
RESPONSE_LEVEL = 0.90
SMOOTH_POINTS = 5
INITIAL_SLOPE_WINDOW_S = 10.0
FINAL_SLOPE_WINDOW_S = 20.0

QUARTZ_DENSITY = 2.648
QUARTZ_SHEAR_MODULUS = 2.947e11

SAVE_EXCEL = True
MAKE_LATEX_TABLE = True
OPEN_RESULTS_FOLDER_AT_END = True
PLOT_DPI = 200


import csv
import glob
import os
import re
import subprocess
import sys
import traceback

AUTO_ONLY = ("--auto" in sys.argv[1:]) or not HAND_SELECTION

try:
    import numpy as np
    import matplotlib
    if AUTO_ONLY:
        matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.patches import Patch
    from matplotlib.widgets import Button
except ImportError as _e:
    print("\nA required package is missing:", _e)
    print("Open a terminal and run:   pip install numpy matplotlib\n")
    sys.exit(1)

if not AUTO_ONLY and matplotlib.get_backend().lower() in ("agg", "pdf", "ps", "svg", "template"):
    print("\n!! Matplotlib cannot open windows here (backend: %s)." % matplotlib.get_backend())
    print("   Run the script from a normal terminal (not from inside a notebook).")
    print("   Continuing with automatic detection only.\n")
    AUTO_ONLY = True

SIGN = 1.0 if DELTA_F_SIGN == "f0_minus_f" else -1.0
DF_LABEL = "Δf = f₀ − f (Hz)" if SIGN > 0 else "Δf = f − f₀ (Hz)"

COLOR_ON = "#f7f06d"
COLOR_OFF = "#b6f08a"
COLOR_OTHER = "#c8c8c8"
COLOR_EXPOSURE = "#dbe9ff"


def read_sqm160_file(path):
    rows, start_info = [], ""
    with open(path, "r", encoding="utf-8", errors="ignore") as fh:
        for line in fh:
            line = line.strip()
            if not line:
                continue
            if line.lower().startswith("start"):
                start_info = " ".join(line.split())
            parts = re.split(r"[\t;]+|\s{2,}|\s", line)
            try:
                rows.append([float(p.replace(",", ".")) for p in parts if p != ""])
            except ValueError:
                continue
    if not rows:
        raise ValueError("No numeric data found in the file.")
    ncol = min(len(r) for r in rows)
    data = np.array([r[:ncol] for r in rows])
    t = data[:, 0]

    if ncol == 2:
        return t, data[:, 1], 1, start_info

    channels = {}
    ch = 1
    while 5 + 3 * (ch - 1) < ncol:
        channels[ch] = 5 + 3 * (ch - 1)
        ch += 1
    if not channels:
        raise ValueError("Could not find a frequency column.")
    if FREQ_CHANNEL == "auto":
        chosen = next((c for c, col in channels.items()
                       if np.median(data[:, col]) > 1000), None)
        if chosen is None:
            raise ValueError("No channel with a valid frequency (> 1 kHz) found.")
    else:
        chosen = int(FREQ_CHANNEL)
        if chosen not in channels:
            raise ValueError(f"Channel {chosen} is not in this file.")
    f = data[:, channels[chosen]]

    keep = f > 1000
    return t[keep], f[keep], chosen, start_info


def labels_from_filename(path):
    name = os.path.splitext(os.path.basename(path))[0].lower()
    tokens = re.split(r"[_\-\s]+", name)

    coating = "?"
    if "4vp" in name:
        coating = "4VP"
    elif "v4d4" in name:
        coating = "pV4D4"
    if coating != "?" and "dvb" in tokens:
        coating += "-DVB"
    elif coating == "?" and "dvb" in tokens:
        coating = "DVB"

    voc_map = {"ace": "Acetone", "acetone": "Acetone", "ipa": "IPA",
               "isopropanol": "IPA", "meth": "Methanol", "methanol": "Methanol",
               "meoh": "Methanol", "eth": "Ethanol", "ethanol": "Ethanol",
               "etoh": "Ethanol", "water": "Water", "h2o": "Water",
               "tol": "Toluene", "toluene": "Toluene"}
    voc = next((voc_map[tk] for tk in tokens if tk in voc_map), "?")

    m = re.search(r"(\d+(?:\.\d+)?)\s*ppm", name)
    conc = float(m.group(1)) if m else None
    m = re.search(r"(\d+(?:\.\d+)?)\s*(?:rh|%rh)", name)
    rh = float(m.group(1)) if m else None
    return coating, voc, conc, rh


def moving_median(y, n):
    if n <= 1:
        return y.copy()
    half = n // 2
    padded = np.pad(y, half, mode="edge")
    return np.array([np.median(padded[i:i + n]) for i in range(len(y))])


def median_in(t, y, t0, t1, fallback):
    sel = (t >= t0) & (t <= t1)
    return float(np.median(y[sel])) if sel.any() else float(fallback)


def slope_in(t, y, t0, t1):
    sel = (t >= t0) & (t <= t1)
    if sel.sum() < 3:
        return np.nan
    return float(np.polyfit(t[sel], y[sel], 1)[0])


def noise_level(s):
    if len(s) < 4:
        return 1e-3
    d2 = np.diff(s, 2)
    sigma_point = 1.4826 * np.median(np.abs(d2 - np.median(d2))) / np.sqrt(6)
    return max(sigma_point * np.sqrt(2), 1e-3)


def find_sudden_changes(t, s):
    sigma = noise_level(s)
    thr = max(MIN_JUMP_HZ, SENSITIVITY * sigma)
    d = np.diff(s)
    flagged = set(np.where(np.abs(d) > thr)[0])
    for k, factor in ((2, 1.6), (3, 2.0)):
        if len(s) > k:
            dk = s[k:] - s[:-k]
            flagged |= set(np.where(np.abs(dk) > factor * thr)[0])
    flagged = sorted(flagged)

    groups = []
    for i in flagged:
        if groups and t[i] - t[groups[-1][-1] + 1] <= MERGE_WINDOW_S:
            groups[-1].append(i)
        else:
            groups.append([i])

    events = []
    for g in groups:
        i0, i1 = g[0], min(g[-1] + 2, len(t) - 1)
        t0, t1 = t[i0], t[i1]
        before = median_in(t, s, t0 - 5.0, t0, s[i0])
        after = median_in(t, s, t1, t1 + 5.0, s[i1])
        step = after - before
        later = median_in(t, s, t1 + 5.0, t1 + 12.0, after) - before
        if step < thr and later >= 1.5 * thr and abs(step) < thr:
            step = later
        if step >= thr:
            direction = "up"
        elif step <= -thr:
            direction = "down"
        else:
            direction = "spike"
        events.append({"t_start": float(t0), "t_end": float(t1), "i_start": i0,
                       "i_end": i1, "step": float(step), "direction": direction,
                       "label": "other", "assumed": False})
    return events, sigma, thr


def refine_on_time(t, s, ev, noise):
    t0 = ev["t_start"]
    sm = moving_median(s, 3)
    base = median_in(t, s, t0 - 15.0, t0 - 8.0, np.nan)
    if np.isnan(base):
        return float(t0)
    onset = t0
    for i in np.where((t >= t0 - 10.0) & (t < t0))[0][::-1]:
        if sm[i] <= base + 2 * noise:
            break
        onset = t[i]
    return float(onset)


def label_events(t, s, events, noise):
    if EVENT_LABELING == "alternate":
        for k, ev in enumerate(events):
            ev["label"] = "ON" if k % 2 == 0 else "OFF"
    else:
        expect = "ON"
        for ev in events:
            if expect == "ON" and ev["direction"] == "down" and not any(
                    e["label"] == "ON" for e in events):
                ev["label"], expect = "OFF", "ON"
                ev["needs_start_on"] = True
            elif expect == "ON" and ev["direction"] == "up":
                ev["label"], expect = "ON", "OFF"
            elif expect == "OFF" and ev["direction"] == "down":
                ev["label"], expect = "OFF", "ON"
            else:
                ev["label"] = "other"

    first = next((e for e in events if e["label"] in ("ON", "OFF")), None)
    if first is not None and first["label"] == "OFF":
        events.insert(0, {"t_start": float(t[0]), "t_end": float(t[0]),
                          "i_start": 0, "i_end": 0, "step": np.nan,
                          "direction": "start", "label": "ON", "assumed": True})

    for ev in events:
        if ev["label"] == "ON" and not ev["assumed"]:
            ev["t_on"] = refine_on_time(t, s, ev, noise)
        else:
            ev["t_on"] = ev["t_start"]
    return events


def make_event(t, s, time_s, label, noise):
    i = int(np.argmin(np.abs(t - time_s)))
    i_end = min(i + 2, len(t) - 1)
    before = median_in(t, s, t[i] - 5.0, t[i], s[i])
    after = median_in(t, s, t[i_end], t[i_end] + 5.0, s[i_end])
    ev = {"t_start": float(t[i]), "t_end": float(t[i_end]), "t_on": float(t[i]),
          "i_start": i, "i_end": i_end, "step": float(after - before),
          "direction": "hand", "label": label, "assumed": False}
    if label == "ON" and i > 0:
        ev["t_on"] = refine_on_time(t, s, ev, noise)
    return ev


def events_from_times(t, s, times, noise):
    return [make_event(t, s, tm, "ON" if k % 2 == 0 else "OFF", noise)
            for k, tm in enumerate(sorted(times))]


def times_from_events(events):
    return [e["t_on"] if e["label"] == "ON" else e["t_start"]
            for e in events if e["label"] in ("ON", "OFF")]


def snap_time(t, s, x, noise):
    nearest = int(np.argmin(np.abs(t - x)))
    if not SNAP_TO_JUMP:
        return float(t[nearest])
    idx = np.where(np.abs(t - x) <= SNAP_WINDOW_S)[0]
    idx = idx[idx < len(t) - 1]
    if len(idx) == 0:
        return float(t[nearest])
    jumps = np.abs(s[idx + 1] - s[idx])
    big = np.where(jumps > max(MIN_JUMP_HZ, 5 * noise))[0]
    if len(big):
        return float(t[idx[big[0]]])
    return float(t[nearest])


def load_selection(path):
    times = []
    try:
        with open(path, "r", encoding="utf-8") as fh:
            for line in fh:
                parts = line.split()
                if len(parts) >= 2 and parts[0] in ("ON", "OFF"):
                    times.append(float(parts[1]))
    except (OSError, ValueError):
        return None
    return times or None


def save_selection(path, events):
    with open(path, "w", encoding="utf-8") as fh:
        fh.write("# label\ttime_s\n")
        for e in events:
            fh.write(f"{e['label']}\t{(e['t_on'] if e['label'] == 'ON' else e['t_start']):.2f}\n")


def analyse_cycles(t, s, events, f0, conc):
    sm = moving_median(s, SMOOTH_POINTS)
    cf = 2 * f0 ** 2 / np.sqrt(QUARTZ_DENSITY * QUARTZ_SHEAR_MODULUS) * 1e-9

    cycles = []
    for k, ev in enumerate(events):
        if ev["label"] != "ON":
            continue
        nxt = events[k + 1] if k + 1 < len(events) else None
        if nxt is None or nxt["label"] != "OFF":
            continue
        after_off = events[k + 2] if k + 2 < len(events) else None

        t_on = ev["t_on"]
        t_off = nxt["t_start"]
        t_off_end = nxt["t_end"]
        if after_off is not None:
            t_purge_end = after_off["t_start"]
        elif cycles:
            t_purge_end = min(float(t[-1]), t_off + float(np.median([c["purge_s"] for c in cycles])))
        else:
            t_purge_end = float(t[-1])

        base = median_in(t, s, t_on - BASELINE_WINDOW_S, t_on, s[np.argmin(np.abs(t - t_on))])
        end = median_in(t, s, t_off - END_WINDOW_S, t_off, s[max(nxt["i_start"], 0)])
        exp_sel = (t >= t_on) & (t < t_off)
        peak = float(np.max(sm[exp_sel])) if exp_sel.any() else end
        df_cycle = end - base
        df_peak = peak - base

        t_resp = np.nan
        if df_peak > 0 and exp_sel.any():
            hit = np.where(exp_sel & (sm >= base + RESPONSE_LEVEL * df_peak))[0]
            if len(hit):
                t_resp = float(t[hit[0]] - t_on)

        after_switch = median_in(t, s, t_off_end, t_off_end + 4.0, end)
        purge_end = median_in(t, s, t_purge_end - END_WINDOW_S, t_purge_end, after_switch)
        recovered = end - purge_end
        recovery_pct = 100.0 * recovered / df_cycle if abs(df_cycle) > 1e-9 else np.nan

        yy, tt = s[exp_sel] - base, t[exp_sel]
        auc = float(np.sum((yy[1:] + yy[:-1]) / 2 * np.diff(tt))) if exp_sel.sum() > 1 else np.nan

        cycles.append({
            "cycle": len(cycles) + 1,
            "t_on_s": t_on, "t_off_s": t_off, "t_purge_end_s": t_purge_end,
            "exposure_s": t_off - t_on, "purge_s": t_purge_end - t_off,
            "baseline_Hz": base, "end_of_exposure_Hz": end, "peak_Hz": peak,
            "dF_cycle_Hz": df_cycle, "dF_peak_Hz": df_peak,
            "response_time_s": t_resp,
            "switch_step_Hz": after_switch - end,
            "level_end_of_purge_Hz": purge_end,
            "recovered_Hz": recovered, "recovery_pct": recovery_pct,
            "slope_initial_Hz_per_s": slope_in(t, s, t_on, t_on + INITIAL_SLOPE_WINDOW_S),
            "slope_final_Hz_per_s": slope_in(t, s, t_off - FINAL_SLOPE_WINDOW_S, t_off),
            "area_Hz_s": auc,
            "mass_ng_per_cm2": df_cycle / cf,
            "sensitivity_Hz_per_ppm": df_cycle / conc if conc else np.nan,
        })
    return cycles, cf


HZ_COLUMNS = {"baseline_Hz", "end_of_exposure_Hz", "peak_Hz", "dF_cycle_Hz",
              "dF_peak_Hz", "switch_step_Hz", "level_end_of_purge_Hz",
              "recovered_Hz", "slope_initial_Hz_per_s", "slope_final_Hz_per_s",
              "area_Hz_s", "sensitivity_Hz_per_ppm"}


def signed(row):
    out = dict(row)
    for k in HZ_COLUMNS:
        if k in out and out[k] is not None:
            out[k] = SIGN * out[k]
    return out


def draw_signal(ax, t, s, events, cycles, title, numbered=False):
    y = SIGN * s
    for c in cycles:
        ax.axvspan(c["t_on_s"], c["t_off_s"], color=COLOR_EXPOSURE, alpha=0.6, lw=0, zorder=0)

    used = set()
    for n, ev in enumerate(events, 1):
        color = {"ON": COLOR_ON, "OFF": COLOR_OFF}.get(ev["label"], COLOR_OTHER)
        name = {"ON": "Sudden change: gas ON", "OFF": "Sudden change: gas OFF"}.get(
            ev["label"], "Other sudden change")
        t0 = ev["t_on"] if ev["label"] == "ON" else ev["t_start"]
        t1 = max(ev["t_end"], t0)
        ax.axvspan(t0 - 3.0, t1 + 3.0, color=color, alpha=0.75, lw=0, zorder=1,
                   label=None if name in used else name)
        used.add(name)
        if numbered:
            ax.text(t0, 1.0, f" {n}: {ev['label']}", transform=ax.get_xaxis_transform(),
                    rotation=90, va="top", ha="right", fontsize=8, color="#444", zorder=7)

    ax.plot(t, y, color="#1f77b4", lw=1.1, zorder=3, label="QCM signal")

    for c in cycles:
        yb, ye = SIGN * c["baseline_Hz"], SIGN * c["end_of_exposure_Hz"]
        ax.hlines(yb, c["t_on_s"], c["t_off_s"], colors="k", linestyles=":", lw=0.9, zorder=4)
        ax.annotate("", xy=(c["t_off_s"] - 2, ye), xytext=(c["t_off_s"] - 2, yb),
                    arrowprops=dict(arrowstyle="<->", color="crimson", lw=1.4), zorder=5)
        ax.plot([c["t_on_s"], c["t_off_s"] - 2], [yb, ye], "o", color="crimson", ms=4, zorder=5)
        ax.text(c["t_off_s"] - 5, (yb + ye) / 2,
                f"cycle {c['cycle']}\nΔf = {SIGN * c['dF_cycle_Hz']:.1f} Hz",
                ha="right", va="center", fontsize=9, color="crimson",
                bbox=dict(boxstyle="round,pad=0.2", fc="white", ec="none", alpha=0.8), zorder=6)

    ax.set_title(title)
    ax.set_xlabel("Time (s)")
    ax.set_ylabel(DF_LABEL)
    ax.grid(alpha=0.35)
    handles, _ = ax.get_legend_handles_labels()
    handles.append(Patch(color=COLOR_EXPOSURE, label="Exposure period"))
    ax.legend(handles=handles, loc="best", fontsize=8, framealpha=0.9)


def plot_run(t, s, events, cycles, title, path_png):
    fig, ax = plt.subplots(figsize=(12, 5.5))
    draw_signal(ax, t, s, events, cycles, title)
    fig.tight_layout()
    fig.savefig(path_png, dpi=PLOT_DPI)
    plt.close(fig)


class CycleSelector:

    HELP = ("LEFT click: add a mark at a sudden change     RIGHT click: remove nearest mark\n"
            "Marks are named ON, OFF, ON, OFF ... from left to right.   "
            "Zoom: magnifier in the toolbar (click it again to go back to marking).")

    def __init__(self, t, s, start_times, auto_times, f0, conc, title, noise):
        self.t, self.s, self.f0, self.conc = t, s, f0, conc
        self.title, self.noise = title, noise
        self.times = sorted(start_times)
        self.auto_times = sorted(auto_times)
        self.history = []
        self.first_draw = True

        self.fig = plt.figure(figsize=(13, 7.5))
        try:
            self.fig.canvas.manager.set_window_title("Select the cycles - " + title)
        except Exception:
            pass
        self.ax = self.fig.add_axes([0.06, 0.17, 0.92, 0.66])
        self.fig.text(0.06, 0.975, self.HELP, fontsize=9.5, va="top", family="monospace")
        self.buttons = []
        specs = [("Undo", self.undo), ("Clear all", self.clear),
                 ("Automatic", self.use_auto), ("Done - calculate", self.finish)]
        for k, (label, func) in enumerate(specs):
            bax = self.fig.add_axes([0.06 + k * 0.18, 0.03, 0.16, 0.07])
            btn = Button(bax, label, color="#e8f5e1" if k == 3 else "0.9", hovercolor="0.8")
            btn.on_clicked(func)
            self.buttons.append(btn)
        self.fig.canvas.mpl_connect("button_press_event", self.on_click)
        self.redraw()

    def remember(self):
        self.history.append(list(self.times))

    def on_click(self, event):
        if event.inaxes is not self.ax or event.xdata is None:
            return
        toolbar = getattr(self.fig.canvas, "toolbar", None)
        if toolbar is not None and getattr(toolbar, "mode", ""):
            return
        if event.button == 1:
            self.remember()
            self.times = sorted(self.times + [snap_time(self.t, self.s, event.xdata, self.noise)])
        elif event.button == 3 and self.times:
            self.remember()
            k = int(np.argmin([abs(x - event.xdata) for x in self.times]))
            self.times.pop(k)
        self.redraw()

    def undo(self, _event=None):
        if self.history:
            self.times = self.history.pop()
            self.redraw()

    def clear(self, _event=None):
        self.remember()
        self.times = []
        self.redraw()

    def use_auto(self, _event=None):
        self.remember()
        self.times = list(self.auto_times)
        self.redraw()

    def finish(self, _event=None):
        plt.close(self.fig)

    def redraw(self):
        if not self.first_draw:
            xlim, ylim = self.ax.get_xlim(), self.ax.get_ylim()
        self.ax.clear()
        events = events_from_times(self.t, self.s, self.times, self.noise)
        cycles, _ = analyse_cycles(self.t, self.s, events, self.f0, self.conc)
        if cycles:
            dfs = ", ".join(f"{SIGN * c['dF_cycle_Hz']:.1f}" for c in cycles)
            info = f"{self.title}   -   {len(cycles)} cycle(s):  Δf = {dfs} Hz"
        else:
            info = f"{self.title}   -   no complete cycle yet (click gas ON, then gas OFF)"
        if len(self.times) % 2 == 1:
            info += "   [last mark has no OFF: used only as end of the last purge]"
        draw_signal(self.ax, self.t, self.s, events, cycles, info, numbered=True)
        if not self.first_draw:
            self.ax.set_xlim(xlim)
            self.ax.set_ylim(ylim)
        self.first_draw = False
        self.fig.canvas.draw_idle()

    def run(self):
        plt.show(block=True)
        return events_from_times(self.t, self.s, self.times, self.noise)


def plot_overlay(t, s, cycles, title, path_png):
    if not cycles:
        return
    fig, ax = plt.subplots(figsize=(8, 5))
    cmap = plt.get_cmap("viridis")
    for k, c in enumerate(cycles):
        sel = (t >= c["t_on_s"]) & (t <= c["t_purge_end_s"])
        col = cmap(k / max(len(cycles) - 1, 1))
        ax.plot(t[sel] - c["t_on_s"], SIGN * (s[sel] - c["baseline_Hz"]), color=col, lw=1.3,
                label=f"cycle {c['cycle']}  (Δf = {SIGN * c['dF_cycle_Hz']:.1f} Hz)")
        ax.axvline(c["exposure_s"], color=col, ls="--", lw=0.8)
    ax.axhline(0, color="k", lw=0.6)
    ax.set_title(title + "  -  cycles overlaid (dashed = gas OFF)")
    ax.set_xlabel("Time since gas ON (s)")
    ax.set_ylabel(DF_LABEL.replace("(Hz)", "relative to cycle baseline (Hz)"))
    ax.grid(alpha=0.35)
    ax.legend(fontsize=8)
    fig.tight_layout()
    fig.savefig(path_png, dpi=PLOT_DPI)
    plt.close(fig)


def plot_comparison(summary, path_png):
    if len(summary) < 2:
        return
    names = [r["label"] for r in summary]
    means = [r["dF_mean_Hz"] for r in summary]
    sds = [0 if np.isnan(r["dF_std_Hz"]) else r["dF_std_Hz"] for r in summary]
    fig, ax = plt.subplots(figsize=(max(6, 1.3 * len(names)), 5))
    ax.bar(range(len(names)), means, yerr=sds, capsize=5, color="#5b8fd1", edgecolor="k")
    ax.set_xticks(range(len(names)))
    ax.set_xticklabels(names, rotation=30, ha="right", fontsize=9)
    ax.set_ylabel("Mean Δf per cycle (Hz)   [error bar = std over cycles]")
    ax.set_title("Comparison of all measurements")
    ax.grid(axis="y", alpha=0.35)
    fig.tight_layout()
    fig.savefig(path_png, dpi=PLOT_DPI)
    plt.close(fig)


def fmt(v, nd=3):
    if v is None:
        return ""
    if isinstance(v, float):
        return "" if np.isnan(v) else f"{v:.{nd}f}"
    return str(v)


def write_csv(path, rows, columns):
    with open(path, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh)
        w.writerow(columns)
        for r in rows:
            w.writerow([fmt(r.get(c)) for c in columns])


def write_excel(path, sheets):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError:
        return False
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows, cols in sheets:
        ws = wb.create_sheet(name[:31])
        ws.append(cols)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for r in rows:
            vals = []
            for c in cols:
                v = r.get(c)
                if isinstance(v, float) and np.isnan(v):
                    v = None
                elif isinstance(v, (np.floating,)):
                    v = float(v)
                vals.append(round(v, 4) if isinstance(v, float) else v)
            ws.append(vals)
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = max(12, len(str(col[0].value)) + 2)
    wb.save(path)
    return True


def write_latex_table(path, summary):
    lines = [
        r"% Generated by qcm_cycle_analyzer.py  -  needs \usepackage{booktabs}",
        r"\begin{table}[h]",
        r"\centering",
        r"\caption{Per-cycle response of polymer-coated QCM sensors (mean $\pm$ standard deviation over exposure cycles).}",
        r"\begin{tabular}{llcccc}",
        r"\toprule",
        r"\textbf{Coating} & \textbf{VOC} & \textbf{Conc. (ppm)} & \textbf{$\Delta f$ per cycle (Hz)} & \textbf{$t_{90}$ (s)} & \textbf{Recovery (\%)} \\",
        r"\midrule",
    ]
    for r in summary:
        conc = "" if r["concentration_ppm"] is None else f"{r['concentration_ppm']:g}"
        def pm(m, sd, nd=1):
            if np.isnan(m):
                return "--"
            return f"{m:.{nd}f}" if np.isnan(sd) else f"{m:.{nd}f} $\\pm$ {sd:.{nd}f}"
        lines.append(f"{r['coating']} & {r['voc']} & {conc} & "
                     f"{pm(r['dF_mean_Hz'], r['dF_std_Hz'])} & "
                     f"{pm(r['t90_mean_s'], r['t90_std_s'], 0)} & "
                     f"{pm(r['recovery_mean_pct'], r['recovery_std_pct'], 0)} \\\\")
    lines += [r"\bottomrule", r"\end{tabular}", r"\end{table}"]
    with open(path, "w", encoding="utf-8") as fh:
        fh.write("\n".join(lines) + "\n")


CYCLE_COLUMNS = ["cycle", "t_on_s", "t_off_s", "exposure_s", "purge_s",
                 "baseline_Hz", "end_of_exposure_Hz", "dF_cycle_Hz", "dF_peak_Hz",
                 "response_time_s", "switch_step_Hz", "recovered_Hz", "recovery_pct",
                 "slope_initial_Hz_per_s", "slope_final_Hz_per_s", "area_Hz_s",
                 "mass_ng_per_cm2", "sensitivity_Hz_per_ppm"]

EVENT_COLUMNS = ["number", "label", "time_s", "jump_Hz", "direction"]


def mean_sd(values):
    v = np.array([x for x in values if x is not None and not np.isnan(x)], float)
    if len(v) == 0:
        return np.nan, np.nan
    return float(v.mean()), (float(v.std(ddof=1)) if len(v) > 1 else np.nan)


def process_file(path, out_dir):
    base = os.path.splitext(os.path.basename(path))[0]
    coating, voc, conc, rh = labels_from_filename(path)
    t, f, channel, start_info = read_sqm160_file(path)
    f0 = float(f[0])
    s = f0 - f

    title_parts = [p for p in (coating, voc, f"{conc:g} ppm" if conc else None,
                               f"{rh:g} % RH" if rh is not None else None) if p and p != "?"]
    title = " / ".join(title_parts) if title_parts else base

    events, sigma, thr = find_sudden_changes(t, s)
    events = label_events(t, s, events, sigma / np.sqrt(2))
    mode = "automatic"

    selection_file = os.path.join(out_dir, base + "_selection.txt")
    if not AUTO_ONLY:
        auto_times = times_from_events(events)
        saved = load_selection(selection_file) if REUSE_SAVED_SELECTION else None
        if saved:
            print(f"   Starting from your saved marks ({os.path.basename(selection_file)}).")
        print("   A window is open: check the marks, then press 'Done - calculate'.")
        selector = CycleSelector(t, s, saved or auto_times, auto_times, f0, conc,
                                 title, sigma / np.sqrt(2))
        events = selector.run()
        save_selection(selection_file, events)
        mode = "by hand"

    cycles, cf = analyse_cycles(t, s, events, f0, conc)

    plot_run(t, s, events, cycles, title, os.path.join(out_dir, base + "_cycles.png"))
    plot_overlay(t, s, cycles, title, os.path.join(out_dir, base + "_overlay.png"))

    signed_cycles = [signed(c) for c in cycles]
    write_csv(os.path.join(out_dir, base + "_cycles.csv"), signed_cycles, CYCLE_COLUMNS)

    event_rows = []
    for n, ev in enumerate(events, 1):
        event_rows.append({"number": n, "label": ev["label"] + (" (assumed at start)" if ev["assumed"] else ""),
                           "time_s": ev["t_on"] if ev["label"] == "ON" else ev["t_start"],
                           "jump_Hz": SIGN * ev["step"],
                           "direction": ev["direction"]})
    write_csv(os.path.join(out_dir, base + "_events.csv"), event_rows, EVENT_COLUMNS)
    if SAVE_EXCEL:
        write_excel(os.path.join(out_dir, base + "_results.xlsx"),
                    [("Cycles", signed_cycles, CYCLE_COLUMNS),
                     ("Sudden changes", event_rows, EVENT_COLUMNS)])

    noise_rms = sigma / np.sqrt(2)
    dF_m, dF_sd = mean_sd([c["dF_cycle_Hz"] for c in cycles])
    t90_m, t90_sd = mean_sd([c["response_time_s"] for c in cycles])
    rec_m, rec_sd = mean_sd([c["recovery_pct"] for c in cycles])
    sens = dF_m / conc if conc and not np.isnan(dF_m) else np.nan
    lod = 3 * noise_rms / sens if sens and not np.isnan(sens) and sens > 0 else np.nan
    sm = moving_median(s, SMOOTH_POINTS)
    baselines = [c["baseline_Hz"] for c in cycles]
    drift = (baselines[-1] - baselines[0]) / (len(baselines) - 1) if len(baselines) > 1 else np.nan

    summary = {
        "file": os.path.basename(path), "label": title, "coating": coating, "voc": voc,
        "concentration_ppm": conc, "humidity_RH": rh, "channel": channel,
        "f0_Hz": f0, "n_cycles": len(cycles),
        "dF_mean_Hz": SIGN * dF_m if not np.isnan(dF_m) else np.nan, "dF_std_Hz": dF_sd,
        "t90_mean_s": t90_m, "t90_std_s": t90_sd,
        "recovery_mean_pct": rec_m, "recovery_std_pct": rec_sd,
        "repeatability_CV_pct": 100 * dF_sd / abs(dF_m) if dF_m and not np.isnan(dF_sd) else np.nan,
        "baseline_drift_Hz_per_cycle": SIGN * drift if not np.isnan(drift) else np.nan,
        "max_shift_from_start_Hz": SIGN * float(np.max(sm)),
        "noise_rms_Hz": noise_rms,
        "sensitivity_Hz_per_ppm": SIGN * sens if not np.isnan(sens) else np.nan,
        "LOD_estimate_ppm": lod,
        "mass_mean_ng_per_cm2": dF_m / cf if not np.isnan(dF_m) else np.nan,
    }

    R = []
    R.append(f"QCM cycle analysis - {os.path.basename(path)}")
    R.append("=" * 70)
    if start_info:
        R.append(start_info)
    R.append(f"Coating: {coating}   VOC: {voc}   Concentration: {conc if conc else '?'} ppm"
             + (f"   Humidity: {rh:g} % RH" if rh is not None else ""))
    R.append(f"Channel {channel}, f0 = {f0:.1f} Hz, {len(t)} points, {t[0]:.1f}-{t[-1]:.1f} s")
    R.append(f"Noise (rms): {noise_rms:.3f} Hz   Jump threshold: "
             + (f"{thr:.2f} Hz" if not np.isnan(thr) else "manual times used")
             + f"   Event naming: {mode}")
    R.append(f"Sign convention: {DF_LABEL}")
    R.append("")
    R.append("SUDDEN CHANGES")
    for e in event_rows:
        R.append(f"  {e['number']:>2}. {e['label']:<26} t = {e['time_s']:7.1f} s   "
                 f"jump = {fmt(e['jump_Hz'], 2):>7} Hz")
    R.append("")
    R.append("CYCLES")
    R.append(f"  {'#':>2} {'ON (s)':>8} {'OFF (s)':>8} {'dF (Hz)':>9} {'peak':>7} "
             f"{'t90 (s)':>8} {'switch':>8} {'recov.%':>8} {'mass ng/cm2':>12}")
    for c in signed_cycles:
        R.append(f"  {c['cycle']:>2} {c['t_on_s']:8.1f} {c['t_off_s']:8.1f} "
                 f"{c['dF_cycle_Hz']:9.2f} {c['dF_peak_Hz']:7.2f} {fmt(c['response_time_s'], 1):>8} "
                 f"{c['switch_step_Hz']:8.2f} {fmt(c['recovery_pct'], 0):>8} "
                 f"{c['mass_ng_per_cm2']:12.1f}")
    if not cycles:
        R.append("  !! No complete gas ON -> gas OFF cycle was found in this file.")
        R.append("     Run the script again (without --auto) and click gas ON and gas OFF")
        R.append("     for every cycle in the window.")
    R.append("")
    R.append("SUMMARY")
    R.append(f"  Mean dF per cycle      : {fmt(summary['dF_mean_Hz'], 2)} +/- {fmt(dF_sd, 2)} Hz"
             f"   (repeatability CV = {fmt(summary['repeatability_CV_pct'], 1)} %)")
    R.append(f"  Max shift from t = 0   : {fmt(summary['max_shift_from_start_Hz'], 2)} Hz"
             "   (includes baseline drift - this is how Table 3 was estimated)")
    R.append(f"  Baseline drift         : {fmt(summary['baseline_drift_Hz_per_cycle'], 2)} Hz per cycle")
    R.append(f"  Mean response time t90 : {fmt(t90_m, 1)} s")
    R.append(f"  Mean recovery          : {fmt(rec_m, 0)} %")
    if conc:
        R.append(f"  Sensitivity            : {fmt(summary['sensitivity_Hz_per_ppm'], 4)} Hz/ppm")
        R.append(f"  LOD estimate (3*noise) : {fmt(lod, 0)} ppm   (assumes a linear response - "
                 "check with several concentrations)")
    R.append("")
    R.append("Column meanings: see README.md")
    with open(os.path.join(out_dir, base + "_report.txt"), "w", encoding="utf-8") as fh:
        fh.write("\n".join(R) + "\n")

    all_rows = []
    for c in signed_cycles:
        row = {"file": os.path.basename(path), "coating": coating, "voc": voc,
               "concentration_ppm": conc, "humidity_RH": rh}
        row.update(c)
        all_rows.append(row)
    return summary, all_rows, R


def files_from_command_line():
    files, missing = [], []
    for arg in sys.argv[1:]:
        if arg.startswith("--"):
            continue
        matches = sorted(glob.glob(arg))
        if matches:
            files.extend(m for m in matches if os.path.isfile(m))
        else:
            missing.append(arg)
    for m in missing:
        print(f"!! File not found: {m}")
    return files


USAGE = """
Usage:
    python qcm_cycle_analyzer.py <data file(s).txt> [--auto]

Examples:
    python qcm_cycle_analyzer.py 4vp_dvb_ace_1000ppm.txt
    python qcm_cycle_analyzer.py file1.txt file2.txt
    python qcm_cycle_analyzer.py *.txt
    python qcm_cycle_analyzer.py *.txt --auto      (no windows, automatic only)
"""


def open_folder(path):
    try:
        if sys.platform.startswith("win"):
            os.startfile(path)
        elif sys.platform == "darwin":
            subprocess.Popen(["open", path])
        else:
            subprocess.Popen(["xdg-open", path])
    except Exception:
        pass


SUMMARY_COLUMNS = ["file", "coating", "voc", "concentration_ppm", "humidity_RH", "n_cycles",
                   "dF_mean_Hz", "dF_std_Hz", "repeatability_CV_pct", "t90_mean_s", "t90_std_s",
                   "recovery_mean_pct", "recovery_std_pct", "baseline_drift_Hz_per_cycle",
                   "max_shift_from_start_Hz", "noise_rms_Hz", "sensitivity_Hz_per_ppm",
                   "LOD_estimate_ppm", "mass_mean_ng_per_cm2", "f0_Hz", "channel"]


def main():
    files = files_from_command_line()
    if not files:
        print(USAGE)
        sys.exit(1)
    out_dir = OUTPUT_FOLDER or os.path.join(os.path.dirname(os.path.abspath(files[0])), "QCM_results")
    os.makedirs(out_dir, exist_ok=True)

    summaries, all_cycles, failed = [], [], []
    for path in files:
        print(f"\n>>> {os.path.basename(path)}")
        try:
            summary, rows, report = process_file(path, out_dir)
            summaries.append(summary)
            all_cycles.extend(rows)
            print("\n".join(report[1:-1]))
        except Exception as e:
            failed.append(path)
            print(f"   !! Could not analyse this file: {e}")
            traceback.print_exc()

    if summaries:
        all_cols = ["file", "coating", "voc", "concentration_ppm", "humidity_RH"] + CYCLE_COLUMNS
        write_csv(os.path.join(out_dir, "ALL_cycles.csv"), all_cycles, all_cols)
        write_csv(os.path.join(out_dir, "ALL_summary.csv"), summaries, SUMMARY_COLUMNS)
        excel_ok = False
        if SAVE_EXCEL:
            excel_ok = write_excel(os.path.join(out_dir, "ALL_results.xlsx"),
                                   [("Summary", summaries, SUMMARY_COLUMNS),
                                    ("All cycles (ML dataset)", all_cycles, all_cols)])
            if not excel_ok:
                print("\n(Excel files skipped - run  pip install openpyxl  to get them.)")
        if MAKE_LATEX_TABLE:
            write_latex_table(os.path.join(out_dir, "ALL_table.tex"), summaries)
        plot_comparison(summaries, os.path.join(out_dir, "ALL_comparison.png"))

    print("\n" + "=" * 70)
    print(f"Done. {len(summaries)} file(s) analysed, {len(failed)} failed.")
    print(f"Results saved in: {out_dir}")
    if OPEN_RESULTS_FOLDER_AT_END and summaries:
        open_folder(out_dir)


if __name__ == "__main__":
    main()
